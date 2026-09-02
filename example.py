"""
Build a three-question ClassPoint deck from scratch.

    uv run --with python-pptx example.py path/to/a-deck-with-classpoint-on-it.pptx
"""

import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu, Pt

import classpoint as cp

OUT = Path('example-out.pptx')
QUESTIONS = [
    ('WordCloud', 'In one word: how should it feel?', cp.word_cloud(submissions=5)),
    ('MultipleChoice', 'Which one is closest?', cp.multiple_choice(['A', 'B', 'C', 'D'])),
    ('ShortAnswers', 'Name one thing it could be built from.', cp.short_answer()),
]


def main(donor):
    # The buttons are Inknoe's artwork. Lift them out of a deck you already have.
    found = cp.extract_buttons(donor)
    print('extracted:', ', '.join(found) or 'nothing — does that deck have activities on it?')

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)

    for index, (kind, prompt, _) in enumerate(QUESTIONS, start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Emu(838200), Emu(1200000), Emu(9800000), Emu(1600000))
        run = box.text_frame.paragraphs[0].add_run()
        run.text = prompt
        run.font.size = Pt(40)
        run.font.bold = True

        # The button must exist before inject() can bind an activity to it, and
        # it must carry ClassPoint's own shape name.
        slide.shapes.add_picture(
            str(cp.BUTTON_IMAGE[kind]),
            Emu(12192000 - 838200 - cp.BUTTON_W_EMU),
            Emu(6858000 - 1000000 - cp.BUTTON_H_EMU),
            Emu(cp.BUTTON_W_EMU), Emu(cp.BUTTON_H_EMU),
        ).name = cp.SHAPE_NAME

    prs.save(OUT)
    cp.inject(OUT, {i: activity for i, (_, _, activity) in enumerate(QUESTIONS, start=1)})

    print(f'wrote {OUT}')
    for row in cp.verify(OUT):
        print(f"  slide {row['slide']}: {row['name']:<15} bound={row['bound']}")


if __name__ == '__main__':
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    main(Path(sys.argv[1]))
